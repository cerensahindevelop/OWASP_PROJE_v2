"""Kurumsal terim sozlugu dosyalarini (.txt/.csv/.xlsx) BELLEKTE
ayristiran, guvenlik sinirlari uygulayan saf modul.

Diske hicbir sekilde yazilmaz - tum girdi/cikti bellek-ici (bytes/BytesIO/
str). Bu dosyalar tanim geregi kurumun en hassas envanteri oldugu icin:
  - Hata mesajlari SADECE konum tasir (sayfa/satir), asla hucre/satir
    ICERIGI - syntax_validator.py'deki "ham icerik hata mesajina
    sizdirilmaz" ilkesiyle ayni.
  - .xlsx bir ZIP arsivi oldugu icin openpyxl'e vermeden ONCE zip
    girislerini (deklare edilen acilmamis boyut + giris sayisi) inceleyip
    zip-bomb koruma sinirlarini uygular - gercekten inflate etmeden.
  - defusedxml.defuse_stdlib() ile (process basina bir kez, bu modul
    import edildiginde) XML entity expansion / "billion laughs"
    saldirilarina karsi openpyxl'in kullandigi stdlib XML ayristiricilari
    sertlestirilir.
  - .xlsm (makrolu) hem uzanti listesine hic girmeyerek (asagida
    _SUPPORTED_EXTENSIONS) hem de ZIP icerigine (xl/vbaProject.bin - bir
    .xlsm dosyasi .xlsx olarak yeniden adlandirilmis olsa bile) bakilarak
    reddedilir.
"""

from __future__ import annotations

import csv
import io
import zipfile
from pathlib import Path

import defusedxml

# ONEMLI: defuse_stdlib() openpyxl import edilmeden ONCE cagrilmali ki
# openpyxl'in kullandigi xml.etree.ElementTree/xml.sax stdlib modulleri
# entity-expansion'a karsi sertlestirilmis halleriyle yuklensin.
defusedxml.defuse_stdlib()

import openpyxl  # noqa: E402


MAX_UPLOAD_BYTES = 10 * 1024 * 1024  # 10MB - tipik kurumsal sozluk boyutunun cok uzerinde
MAX_XLSX_DECOMPRESSED_BYTES = 50 * 1024 * 1024  # 50MB - zip-bomb sinirini
MAX_XLSX_ZIP_ENTRIES = 2_000  # normal bir .xlsx onlarca giris icerir

_SUPPORTED_EXTENSIONS = {".txt", ".csv", ".xlsx"}
_HEADER_KEYWORDS = {"terim", "term", "kelime", "value", "deger", "değer"}

# Kurum-ici sablon sozlesmesi: .csv/.xlsx dosyasinda bu basliga sahip bir
# sutun varsa (orn. "ID | Kategori | Hassas Deger" gibi cok sutunlu bir
# tablo), SADECE bu sutun terim kaynagi olarak okunur - ID/Kategori gibi
# diger sutunlar sessizce yok sayilir. Bulunamazsa eski tek-sutun (ilk
# sutun) davranisina geri donulur (bkz. _find_value_column_index).
_VALUE_COLUMN_HEADERS = {"hassas deger", "hassas değer"}


class TermFileError(ValueError):
    """Kurumsal terim dosyasi islenirken olusan tum hatalarin temel sinifi."""


class TermFileFormatError(TermFileError):
    """Desteklenmeyen ya da taninmayan dosya formati - sessizce degil,
    acikca reddedilir."""


class TermFileSecurityError(TermFileError):
    """Dosya boyutu/zip-yapisi guvenlik sinirlarindan biri asildi
    (zip-bomb, entity-expansion, makro icerme vb.)."""


class TermFileParseError(TermFileError):
    """Dosya ayristirilamadi. Mesaj SADECE konum icerir (sayfa/satir),
    hucre/satir ICERIGI ASLA yer almaz."""


# Bu modulun tek genel giris noktasi: dosya adindan formati belirler,
# guvenlik sinirlarini uygular, terim listesini dondurur. Bos/whitespace
# satirlar ve (varsa) baslik satiri atlanir.
def parse_terms_from_file(filename: str, content: bytes) -> list[str]:
    if len(content) > MAX_UPLOAD_BYTES:
        raise TermFileSecurityError(
            f"dosya boyutu {MAX_UPLOAD_BYTES} bayt sinirini asiyor ({len(content)} bayt)"
        )

    suffix = Path(filename).suffix.lower()
    if suffix not in _SUPPORTED_EXTENSIONS:
        raise TermFileFormatError(
            f"desteklenmeyen dosya formati: '{suffix or '(uzantisiz)'}' - "
            "sadece .txt, .csv, .xlsx destekleniyor"
        )

    if suffix == ".txt":
        return _parse_txt(content)
    if suffix == ".csv":
        return _parse_csv(content)
    return _parse_xlsx(content)


# Bayt icerigini UTF-8 (BOM'lu olabilir) olarak decode eder; basarisiz
# olursa mesajda SADECE bayt konumunu tasiyan bir hata firlatir.
def _decode_text(content: bytes, *, format_label: str) -> str:
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise TermFileParseError(
            f"{format_label} UTF-8 olarak decode edilemedi (konum: bayt {exc.start})"
        ) from exc


# .txt dosyasini satir satir okur, bos satirlari atlar.
def _parse_txt(content: bytes) -> list[str]:
    text = _decode_text(content, format_label="metin dosyasi")
    return [line.strip() for line in text.splitlines() if line.strip()]


# Baslik satirindaki ham degerler icinde "Hassas Deger" sablonuyla eslesen
# sutunun (0-tabanli) indeksini bulur; yoksa None doner (bkz.
# _VALUE_COLUMN_HEADERS dokstringi - tek-sutun eski davranisa donus sinyali).
def _find_value_column_index(header_values: list) -> int | None:
    for idx, value in enumerate(header_values):
        if value is None:
            continue
        text = str(value).strip().lower()
        if text in _VALUE_COLUMN_HEADERS:
            return idx
    return None


# "Hassas Deger" basligina sahip bir sutun varsa SADECE o sutunu okur (diger
# sutunlar - ID, Kategori vb. - yok sayilir); yoksa ilk kolonu terim listesi
# olarak okur ve ilk satir bilinen bir baslik kelimesiyle
# (terim/term/kelime/deger vb.) eslesirse atlar.
def _parse_csv(content: bytes) -> list[str]:
    text = _decode_text(content, format_label="CSV")
    reader = csv.reader(io.StringIO(text))
    try:
        rows = list(reader)
    except csv.Error as exc:
        raise TermFileParseError(f"CSV ayristirma hatasi (satir {reader.line_num})") from exc

    if not rows:
        return []

    terms: list[str] = []
    value_col = _find_value_column_index(rows[0])
    if value_col is not None:
        for row in rows[1:]:
            if value_col >= len(row):
                continue
            value = row[value_col].strip()
            if value:
                terms.append(value)
        return terms

    for line_no, row in enumerate(rows, start=1):
        if not row:
            continue
        value = row[0].strip()
        if not value:
            continue
        if line_no == 1 and value.lower() in _HEADER_KEYWORDS:
            continue
        terms.append(value)
    return terms


# openpyxl'e vermeden ONCE zip girislerini inceler - gercekten inflate
# etmeden, sadece ZIP merkezi dizinindeki DEKLARE EDILEN acilmamis boyuta
# bakarak zip-bomb'u yakalar (klasik/standart savunma).
def _inspect_xlsx_zip_safety(content: bytes) -> None:
    try:
        zf = zipfile.ZipFile(io.BytesIO(content))
    except zipfile.BadZipFile as exc:
        raise TermFileParseError("gecerli bir .xlsx (zip) dosyasi degil") from exc

    with zf:
        infolist = zf.infolist()
        if len(infolist) > MAX_XLSX_ZIP_ENTRIES:
            raise TermFileSecurityError(
                f".xlsx icindeki dosya sayisi {MAX_XLSX_ZIP_ENTRIES} sinirini asiyor ({len(infolist)})"
            )

        total_declared = 0
        for info in infolist:
            if info.filename == "xl/vbaProject.bin":
                raise TermFileSecurityError(
                    "makro iceren (.xlsm tipi) dosyalar kabul edilmiyor"
                )
            if info.file_size > MAX_XLSX_DECOMPRESSED_BYTES:
                raise TermFileSecurityError(
                    f".xlsx icindeki bir oge tek basina {MAX_XLSX_DECOMPRESSED_BYTES} bayt "
                    "acilmamis-boyut sinirini asiyor"
                )
            total_declared += info.file_size
        if total_declared > MAX_XLSX_DECOMPRESSED_BYTES:
            raise TermFileSecurityError(
                f".xlsx toplam acilmamis boyutu {MAX_XLSX_DECOMPRESSED_BYTES} bayt sinirini asiyor"
            )


# .xlsx'in ilk sayfasinda "Hassas Deger" basligina sahip bir sutun varsa
# SADECE o sutunu okur (diger sutunlar - ID, Kategori vb. - yok sayilir);
# yoksa ilk kolonu terim listesi olarak okur ve ilk satir bir baslik
# kelimesiyse atlar (once zip-guvenlik kontrolunden gecirir).
def _parse_xlsx(content: bytes) -> list[str]:
    _inspect_xlsx_zip_safety(content)

    try:
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl bozuk dosyalarda cesitli exception turleri firlatabilir
        raise TermFileParseError("xlsx dosyasi acilamadi (bozuk ya da desteklenmeyen bicim)") from exc

    try:
        sheet = workbook.active
        if sheet is None:
            raise TermFileParseError("xlsx dosyasinda okunabilir bir sayfa bulunamadi")

        terms: list[str] = []
        row_idx = 0
        try:
            header_row = next(sheet.iter_rows(min_row=1, max_row=1), None)
            header_values = [cell.value for cell in header_row] if header_row else []
            value_col = _find_value_column_index(header_values)

            if value_col is not None:
                col_number = value_col + 1
                for row_idx, row in enumerate(
                    sheet.iter_rows(min_row=2, min_col=col_number, max_col=col_number), start=2
                ):
                    value = row[0].value
                    if value is None:
                        continue
                    text_value = str(value).strip()
                    if text_value:
                        terms.append(text_value)
            else:
                for row_idx, row in enumerate(sheet.iter_rows(min_col=1, max_col=1), start=1):
                    value = row[0].value
                    if value is None:
                        continue
                    text_value = str(value).strip()
                    if not text_value:
                        continue
                    if row_idx == 1 and text_value.lower() in _HEADER_KEYWORDS:
                        continue
                    terms.append(text_value)
        except TermFileParseError:
            raise
        except Exception as exc:
            raise TermFileParseError(
                f"xlsx ayristirma hatasi (konum: {sheet.title}!satir {row_idx})"
            ) from exc
        return terms
    finally:
        workbook.close()
