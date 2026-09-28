"""Kurumsal terim sozlugu ozelligi / Adim 2: dosya ayristirma + guvenlik
sertlestirmesi testleri. DB'siz - tamamen bellek-ici bytes uzerinde calisir.
"""

from __future__ import annotations

import io
import zipfile

import openpyxl
import pytest

from app.services import term_file_parser as parser
from app.services.term_file_parser import (
    TermFileFormatError,
    TermFileParseError,
    TermFileSecurityError,
    parse_terms_from_file,
)

_SENSITIVE_MARKER = "GIZLI-KURUMSAL-DEGER-XYZ123"


def _xlsx_bytes(rows: list[str], *, with_header: bool = False) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    idx = 1
    if with_header:
        ws.cell(row=idx, column=1, value="terim")
        idx += 1
    for value in rows:
        ws.cell(row=idx, column=1, value=value)
        idx += 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------
# Mutlu yol
# --------------------------------------------------------------------------


def test_txt_parses_lines_and_skips_blank():
    content = "Atlas\n\n  Poseidon  \n\nZeus\n".encode("utf-8")
    terms = parse_terms_from_file("liste.txt", content)
    assert terms == ["Atlas", "Poseidon", "Zeus"]


def test_txt_handles_utf8_bom():
    content = "﻿Atlas\nPoseidon\n".encode("utf-8")
    terms = parse_terms_from_file("liste.txt", content)
    assert terms == ["Atlas", "Poseidon"]


def test_csv_parses_first_column_and_skips_header():
    content = "terim\nAtlas\nPoseidon\n".encode("utf-8")
    terms = parse_terms_from_file("liste.csv", content)
    assert terms == ["Atlas", "Poseidon"]


def test_csv_without_header_keeps_all_rows():
    content = "Atlas\nPoseidon\n".encode("utf-8")
    terms = parse_terms_from_file("liste.csv", content)
    assert terms == ["Atlas", "Poseidon"]


def test_xlsx_parses_first_column_and_skips_header():
    content = _xlsx_bytes(["Atlas", "Poseidon"], with_header=True)
    terms = parse_terms_from_file("liste.xlsx", content)
    assert terms == ["Atlas", "Poseidon"]


def test_xlsx_without_header_keeps_all_rows():
    content = _xlsx_bytes(["Atlas", "Poseidon"], with_header=False)
    terms = parse_terms_from_file("liste.xlsx", content)
    assert terms == ["Atlas", "Poseidon"]


def test_xlsx_multi_column_reads_hassas_deger_column_only():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ID", "Kategori", "Hassas Değer"])
    ws.append([1, "Kurum adı", "Atlas"])
    ws.append([2, "Sistem adı", "Poseidon"])
    buf = io.BytesIO()
    wb.save(buf)
    terms = parse_terms_from_file("liste.xlsx", buf.getvalue())
    assert terms == ["Atlas", "Poseidon"]


def test_csv_multi_column_reads_hassas_deger_column_only():
    content = (
        "ID,Kategori,Hassas Değer\n1,Kurum adı,Atlas\n2,Sistem adı,Poseidon\n"
    ).encode("utf-8")
    terms = parse_terms_from_file("liste.csv", content)
    assert terms == ["Atlas", "Poseidon"]


# --------------------------------------------------------------------------
# Format reddi
# --------------------------------------------------------------------------


def test_unsupported_extension_is_rejected_explicitly():
    with pytest.raises(TermFileFormatError, match="desteklenmeyen dosya formati"):
        parse_terms_from_file("liste.docx", b"herhangi bir icerik")


def test_no_extension_is_rejected_explicitly():
    with pytest.raises(TermFileFormatError):
        parse_terms_from_file("liste", b"herhangi bir icerik")


def test_xlsm_extension_is_rejected_without_even_opening():
    with pytest.raises(TermFileFormatError):
        parse_terms_from_file("liste.xlsm", b"herhangi bir icerik")


# --------------------------------------------------------------------------
# Boyut / zip-bomb / entity-expansion sinirlari
# --------------------------------------------------------------------------


def test_oversized_raw_upload_is_rejected(monkeypatch):
    monkeypatch.setattr(parser, "MAX_UPLOAD_BYTES", 10)
    with pytest.raises(TermFileSecurityError, match="dosya boyutu"):
        parse_terms_from_file("liste.txt", b"bu on bayttan uzun bir icerik")


def test_xlsx_too_many_zip_entries_is_rejected(monkeypatch):
    content = _xlsx_bytes(["Atlas"])
    monkeypatch.setattr(parser, "MAX_XLSX_ZIP_ENTRIES", 1)
    with pytest.raises(TermFileSecurityError, match="dosya sayisi"):
        parse_terms_from_file("liste.xlsx", content)


def test_xlsx_declared_decompressed_size_over_limit_is_rejected(monkeypatch):
    content = _xlsx_bytes(["Atlas", "Poseidon", "Zeus"])
    monkeypatch.setattr(parser, "MAX_XLSX_DECOMPRESSED_BYTES", 10)
    with pytest.raises(TermFileSecurityError, match="acilmamis"):
        parse_terms_from_file("liste.xlsx", content)


def test_xlsx_with_macro_entry_is_rejected_even_with_xlsx_extension():
    """Bir .xlsm dosyasi .xlsx olarak yeniden adlandirilmis olsa bile
    (uzanti kontrolunu atlatsa bile) ZIP icerigindeki xl/vbaProject.bin
    varligina bakilarak reddedilmelidir."""
    base = _xlsx_bytes(["Atlas"])
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(base)) as src, zipfile.ZipFile(out, "w") as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.writestr("xl/vbaProject.bin", b"sahte makro icerigi")

    with pytest.raises(TermFileSecurityError, match="makro"):
        parse_terms_from_file("liste.xlsx", out.getvalue())


def test_corrupted_zip_is_reported_as_parse_error_not_crash():
    with pytest.raises(TermFileParseError):
        parse_terms_from_file("liste.xlsx", b"bu gecerli bir zip/xlsx degil")


def test_bad_encoding_txt_is_reported_with_location_only():
    content = b"\xff\xfe\x00bozuk-bayt-dizisi"
    with pytest.raises(TermFileParseError, match="bayt"):
        parse_terms_from_file("liste.txt", content)


# --------------------------------------------------------------------------
# Hata mesajlari icerik SIZDIRMAZ - sadece konum
# --------------------------------------------------------------------------


def test_parse_error_messages_never_contain_cell_content():
    base = _xlsx_bytes(["Atlas"])
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(base)) as src, zipfile.ZipFile(out, "w") as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.writestr("xl/vbaProject.bin", _SENSITIVE_MARKER.encode("utf-8"))

    with pytest.raises(TermFileSecurityError) as excinfo:
        parse_terms_from_file("liste.xlsx", out.getvalue())
    assert _SENSITIVE_MARKER not in str(excinfo.value)


def test_size_limit_error_does_not_echo_file_content(monkeypatch):
    monkeypatch.setattr(parser, "MAX_UPLOAD_BYTES", 5)
    payload = f"{_SENSITIVE_MARKER}\n".encode("utf-8")
    with pytest.raises(TermFileSecurityError) as excinfo:
        parse_terms_from_file("liste.txt", payload)
    assert _SENSITIVE_MARKER not in str(excinfo.value)
